"""Raw text + layout extraction from a resume file.

PDF support (via PyMuPDF) is mandatory and is the only format that yields
span-level layout data, which `extraction.integrity` needs to detect
visually-hidden text. DOCX/TXT are optional, plain-text-only formats.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pymupdf

# page.get_text() clips to the page rect by default; widen it by this much
# so text placed off-page (a hiding technique) is still extracted and can be
# flagged by integrity analysis instead of silently disappearing.
OFF_PAGE_CLIP_MARGIN_PT = 5000


@dataclass
class TextSpan:
    """One PyMuPDF text span: a run of text sharing font/color/position."""

    text: str
    color_rgb: tuple[int, int, int]
    font_size: float
    bbox: tuple[float, float, float, float]
    page_rect: tuple[float, float, float, float]
    page_index: int
    block_index: int
    line_index: int
    opacity: float | None = None  # None when the PDF backend can't report it
    render_mode: int | None = None  # 3 == invisible text per PDF spec


@dataclass
class ExtractedDocument:
    kind: str
    full_text: str
    spans: list[TextSpan] = field(default_factory=list)


def _unpack_color(color_int: int) -> tuple[int, int, int]:
    r = (color_int >> 16) & 0xFF
    g = (color_int >> 8) & 0xFF
    b = color_int & 0xFF
    return r, g, b


def _extract_render_modes(page: pymupdf.Page) -> dict[tuple[float, float, str], tuple[float | None, int | None]]:
    """Best-effort opacity/render-mode lookup keyed by (origin_x, origin_y, text).

    `get_texttrace()` is only available on modern PyMuPDF builds and its
    schema has shifted across versions, so any failure here is swallowed --
    callers just get None back and fall back to color/position heuristics.
    """
    lookup: dict[tuple[float, float, str], tuple[float | None, int | None]] = {}
    try:
        traces = page.get_texttrace()
    except Exception:
        return lookup
    for trace in traces:
        render_mode = trace.get("type")
        opacity = trace.get("opacity")
        for ch in trace.get("chars", []):
            try:
                origin = ch[4] if len(ch) > 4 else None
                text_piece = chr(ch[0]) if isinstance(ch[0], int) else str(ch[0])
            except Exception:
                continue
            if origin is not None:
                key = (round(origin[0], 1), round(origin[1], 1), text_piece)
                lookup[key] = (opacity, render_mode)
    return lookup


def extract_pdf(path: Path) -> ExtractedDocument:
    doc = pymupdf.open(path)
    try:
        full_text_parts: list[str] = []
        spans: list[TextSpan] = []
        for page_index, page in enumerate(doc):
            page_rect = (page.rect.x0, page.rect.y0, page.rect.x1, page.rect.y1)
            # get_text() clips to the page rect by default, which would make
            # genuinely off-page (hidden) text invisible to us too -- widen
            # the clip so such spans are still captured for integrity analysis.
            wide_clip = pymupdf.Rect(
                page.rect.x0 - OFF_PAGE_CLIP_MARGIN_PT,
                page.rect.y0 - OFF_PAGE_CLIP_MARGIN_PT,
                page.rect.x1 + OFF_PAGE_CLIP_MARGIN_PT,
                page.rect.y1 + OFF_PAGE_CLIP_MARGIN_PT,
            )
            full_text_parts.append(page.get_text(clip=wide_clip))
            render_lookup = _extract_render_modes(page)
            raw = page.get_text("dict", clip=wide_clip)
            for block_index, block in enumerate(raw.get("blocks", [])):
                for line_index, line in enumerate(block.get("lines", [])):
                    for span in line.get("spans", []):
                        text = span.get("text", "")
                        if not text.strip():
                            continue
                        color_rgb = _unpack_color(span.get("color", 0))
                        origin = span.get("origin", (span["bbox"][0], span["bbox"][1]))
                        key = (round(origin[0], 1), round(origin[1], 1), text[0])
                        opacity, render_mode = render_lookup.get(key, (None, None))
                        spans.append(
                            TextSpan(
                                text=text,
                                color_rgb=color_rgb,
                                font_size=span.get("size", 0.0),
                                bbox=tuple(span.get("bbox", (0, 0, 0, 0))),
                                page_rect=page_rect,
                                page_index=page_index,
                                block_index=block_index,
                                line_index=line_index,
                                opacity=opacity,
                                render_mode=render_mode,
                            )
                        )
        return ExtractedDocument(kind="pdf", full_text="\n".join(full_text_parts), spans=spans)
    finally:
        doc.close()


def extract_docx(path: Path) -> ExtractedDocument:
    from docx import Document  # local import: optional dependency

    document = Document(str(path))
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)
    return ExtractedDocument(kind="docx", full_text="\n".join(paragraphs))


def extract_txt(path: Path) -> ExtractedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    return ExtractedDocument(kind="txt", full_text=text)


def extract_document(path: Path, kind: str) -> ExtractedDocument:
    if kind == "pdf":
        return extract_pdf(path)
    if kind == "docx":
        return extract_docx(path)
    if kind == "txt":
        return extract_txt(path)
    raise ValueError(f"Unsupported document kind: {kind}")
